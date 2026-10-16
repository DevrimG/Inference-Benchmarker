"""Phase 1: arrival modes, statistics, goodput and the run manifest."""

from __future__ import annotations

import json
import math

import numpy as np
import pytest

from llm_load_tester import stats
from llm_load_tester.arrival import build_schedule
from llm_load_tester.mock_server import MockStreamConfig

from .conftest import run_benchmark

FAST = dict(output_tokens=4, ttft_s=0.01, token_interval_s=0.005)


# -- arrival schedules --------------------------------------------------------

def test_poisson_schedule_is_seeded_and_has_the_requested_rate() -> None:
    a = build_schedule("poisson", requests=4000, rate=10.0, seed=7)
    b = build_schedule("poisson", requests=4000, rate=10.0, seed=7)
    assert a.offsets == b.offsets and a.digest() == b.digest()
    assert a.offsets[0] == 0.0
    gaps = np.diff(a.offsets)
    assert gaps.mean() == pytest.approx(0.1, rel=0.05)
    assert gaps.std() / gaps.mean() == pytest.approx(1.0, rel=0.1)  # exponential: CV = 1


def test_gamma_shape_controls_burstiness() -> None:
    smooth = np.diff(build_schedule("gamma", requests=4000, rate=10, shape=4.0, seed=1).offsets)
    bursty = np.diff(build_schedule("gamma", requests=4000, rate=10, shape=0.25, seed=1).offsets)
    assert smooth.mean() == pytest.approx(0.1, rel=0.05)
    assert bursty.mean() == pytest.approx(0.1, rel=0.15)
    assert smooth.std() / smooth.mean() == pytest.approx(0.5, rel=0.1)  # 1/sqrt(k)
    assert bursty.std() / bursty.mean() == pytest.approx(2.0, rel=0.15)


def test_duration_schedule_covers_duration() -> None:
    schedule = build_schedule("poisson", rate=20, duration_seconds=30, seed=3)
    assert schedule.offsets[-1] <= 30
    assert len(schedule) == pytest.approx(600, rel=0.15)


def test_replay_schedule_round_trips_through_the_log(tmp_path) -> None:
    trace = tmp_path / "trace.jsonl"
    trace.write_text("\n".join(json.dumps(row) for row in [
        {"timestamp": 1500, "fixture_id": "b"},
        {"timestamp": 1000, "fixture_id": "a"},
        {"timestamp": 1750},
    ]))
    schedule = build_schedule("replay", replay_file=trace)
    assert schedule.offsets == [0.0, 0.5, 0.75]
    assert schedule.fixture_ids == ["a", "b", None]
    logged = tmp_path / "logged.jsonl"
    schedule.write_jsonl(logged)
    again = build_schedule("replay", replay_file=logged)
    assert again.offsets == schedule.offsets and again.fixture_ids == schedule.fixture_ids


async def test_open_loop_sends_on_schedule_despite_a_stalled_request(mock_server, tmp_path) -> None:
    """Open loop: a 1 s stall on request 1 must not delay requests 2..6."""
    trace = tmp_path / "trace.jsonl"
    trace.write_text("\n".join(json.dumps({"offset_s": 0.05 * i}) for i in range(6)))
    mock_server.config = MockStreamConfig(**FAST, request_stalls={1: 1.0})
    log = tmp_path / "schedule.jsonl"
    result = await run_benchmark(
        mock_server, arrival_mode="replay", replay_file=str(trace), total_requests=0,
        schedule_log_path=str(log),
    )
    start = result.scenario_start_monotonic
    received = sorted(r.received_at - start for r in mock_server.requests)
    assert len(received) == 6
    for got, want in zip(received, [0.05 * i for i in range(6)]):
        assert got == pytest.approx(want, abs=0.02)
    lags = result.get_summary()["distributions"]["schedule_lag_s"]
    assert lags["n"] == 6 and lags["max"] < 0.02
    assert log.exists() and "schedule" in log.read_text().splitlines()[0]


async def test_burst_releases_exactly_c_requests_together(mock_server) -> None:
    mock_server.config = MockStreamConfig(output_tokens=4, ttft_s=0.1, token_interval_s=0.01)
    result = await run_benchmark(mock_server, arrival_mode="burst", concurrency=8, total_requests=100)
    assert len(result.request_metrics) == 8
    assert mock_server.max_running == 8
    starts = [m.start_time for m in result.request_metrics]
    assert max(starts) - min(starts) < 0.02


async def test_closed_loop_duration_mode_stops_at_deadline(mock_server) -> None:
    mock_server.config = MockStreamConfig(**FAST)
    result = await run_benchmark(mock_server, concurrency=2, duration_seconds=0.5)
    ends = [m.end_time - result.scenario_start_monotonic for m in result.request_metrics]
    assert len(result.request_metrics) > 4
    assert max(ends) < 0.5 + 0.1


# -- statistics ---------------------------------------------------------------

def test_describe_reports_n_and_a_ci_that_covers_the_true_percentile() -> None:
    rng = np.random.default_rng(0)
    covered = 0
    trials = 200
    for trial in range(trials):
        sample = rng.exponential(1.0, size=200)
        summary = stats.describe(sample, (50,), seed=trial)
        lo, hi = summary["ci95"]["p50"]
        covered += lo <= math.log(2) <= hi
    assert summary["n"] == 200 and summary["ci_method"] == "bootstrap"
    assert 0.90 <= covered / trials <= 0.99


def test_large_samples_use_order_statistic_interval() -> None:
    data = np.arange(10_001, dtype=float)
    summary = stats.describe(data, (50, 99))
    assert summary["ci_method"] == "order_statistic"
    lo, hi = summary["ci95"]["p50"]
    assert lo < 5000 < hi and hi - lo < 300


def test_coefficient_of_variation() -> None:
    assert stats.coefficient_of_variation([10, 10, 10]) == 0
    assert stats.coefficient_of_variation([9, 11]) == pytest.approx(math.sqrt(2) / 10)
    assert stats.coefficient_of_variation([1]) is None


def test_steady_state_window_and_littles_law() -> None:
    # Two staggered workers, each running back-to-back 1 s requests.
    intervals = [(0.0, 1.0), (1.0, 2.0), (2.0, 3.0), (0.5, 1.5), (1.5, 2.5)]
    window = stats.steady_state_window(intervals, 2)
    assert window == (0.5, 2.5)
    law = stats.littles_law(intervals, window)
    assert law["observed_mean_in_flight"] == pytest.approx(2.0)
    assert law["predicted_in_flight"] == pytest.approx(2.0)


# -- goodput, summary, manifest -----------------------------------------------

async def test_goodput_against_slo(mock_server) -> None:
    mock_server.config = MockStreamConfig(
        output_tokens=10, ttft_s=0.02, token_interval_s=0.01, request_stalls={3: 0.5}
    )
    result = await run_benchmark(
        mock_server, total_requests=6,
        slo={"ttft_p95_s": 0.2, "decode_p5_tok_s": 50, "max_stall_s": 0.1},
    )
    goodput = result.get_summary()["goodput"]
    assert goodput["good_requests"] == 5
    assert goodput["run_checks"]["decode_p5"]["pass"] is True
    assert goodput["run_checks"]["ttft_p95"]["pass"] is False
    assert goodput["run_slo_pass"] is False


async def test_summary_distributions_carry_n_and_ci(mock_server) -> None:
    mock_server.config = MockStreamConfig(**FAST)
    summary = (await run_benchmark(mock_server, concurrency=2, total_requests=10)).get_summary()
    ttft = summary["distributions"]["ttft_s"]
    assert ttft["n"] == 10 and ttft["ci95"]["p95"] is not None
    assert summary["littles_law"]["available"]
    assert summary["littles_law"]["ratio_predicted_to_observed"] == pytest.approx(1.0, rel=0.2)
    # existing keys are unchanged
    for key in ("ttft_p95", "latency_p99", "tpot_p50", "overall_tokens_per_second"):
        assert key in summary


async def test_manifest_records_provenance(mock_server, tmp_path) -> None:
    mock_server.config = MockStreamConfig(**FAST)
    result = await run_benchmark(
        mock_server, api_key="secret-key", arrival_mode="poisson", arrival_rate=50,
        arrival_seed=11, total_requests=3, image_digest="sha256:abc",
    )
    manifest = result.manifest
    assert manifest["tool"]["sha"]
    assert manifest["served_models"] == ["mock-model"]
    assert manifest["server_version"] == "mock-vllm-0.0.1"
    assert manifest["image_digest"] == "sha256:abc"
    assert manifest["arrival_schedule"]["seed"] == 11
    assert manifest["client_host"]["hostname"]
    assert "secret-key" not in json.dumps(manifest)
    path = tmp_path / "r.json"
    result.export_json(path)
    exported = json.loads(path.read_text())
    assert exported["manifest"]["arrival_schedule"]["sha256"]
    assert exported["raw_metrics"][0]["chunk_records"]
    assert "secret-key" not in path.read_text()


async def test_sent_request_body_is_exported_without_headers(mock_server, tmp_path) -> None:
    from openpyxl import load_workbook

    mock_server.config = MockStreamConfig(**FAST)
    result = await run_benchmark(mock_server, api_key="secret-key", total_requests=2)
    payload = result.request_metrics[0].request_payload
    assert payload == mock_server.requests[0].body  # exactly what the server received
    result.export_json(tmp_path / "r.json")
    result.export_csv(tmp_path / "r.csv")
    result.export_xlsx(tmp_path / "r.xlsx")
    exported = json.loads((tmp_path / "r.json").read_text())["raw_metrics"][0]["request_payload"]
    assert exported["messages"][0]["content"] == "hello"
    assert "request_payload" in (tmp_path / "r.csv").read_text().splitlines()[0]
    sheet = load_workbook(tmp_path / "r.xlsx")["Requests"]
    assert any("hello" in str(c.value) for row in sheet.iter_rows() for c in row)
    for name in ("r.json", "r.csv"):
        assert "secret-key" not in (tmp_path / name).read_text()


def test_media_is_redacted_but_text_is_kept() -> None:
    import base64
    from llm_load_tester.metrics import redact_media

    image = "data:image/png;base64," + base64.b64encode(b"\x89PNG" * 2000).decode()
    body = {"messages": [{"role": "user", "content": [
        {"type": "text", "text": "Describe this image."},
        {"type": "image_url", "image_url": {"url": image}},
    ]}], "max_tokens": 64}
    out = redact_media(body)
    parts = out["messages"][0]["content"]
    assert parts[0]["text"] == "Describe this image." and out["max_tokens"] == 64
    assert parts[1]["image_url"]["url"].startswith("data:image/png;base64,<media sha256=")
    long_text = "word " * 1000
    assert redact_media({"t": long_text})["t"] == long_text  # long prose is not media
