"""Phase 2: fan-out, fixtures, metrics sidecar, ladder, soak, burst, config CLI."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from llm_load_tester.fixtures import FixtureHandler
from llm_load_tester.lab import (
    EndpointSpec,
    LadderRules,
    evaluate_ladder,
    run_fanout,
    soak_windows,
    split_concurrency,
)
from llm_load_tester.lab_cli import run_from_config
from llm_load_tester.mock_server import MockOpenAIServer, MockStreamConfig

from .conftest import make_config, run_benchmark

FAST = dict(output_tokens=6, ttft_s=0.01, token_interval_s=0.005)


@pytest.fixture
async def second_server():
    server = MockOpenAIServer(MockStreamConfig(**FAST))
    await server.start()
    try:
        yield server
    finally:
        await server.stop()


def write_fixtures(path: Path) -> Path:
    path.write_text("\n".join(json.dumps(row) for row in [
        {"fixture_id": "fx-a", "prompt": "alpha", "max_tokens": 8},
        {"fixture_id": "fx-b", "messages": [{"role": "user", "content": "beta"}], "max_tokens": 8},
    ]))
    return path


def test_split_concurrency_largest_remainder() -> None:
    assert split_concurrency(8, [1, 1]) == [4, 4]
    assert split_concurrency(10, [3, 1]) == [8, 2]
    assert split_concurrency(5, [1, 1, 1]) == [2, 2, 1]
    assert sum(split_concurrency(7, [0.2, 0.5, 0.3])) == 7
    with pytest.raises(ValueError):
        split_concurrency(1, [1, 1])


async def test_fanout_shares_start_and_merges_with_endpoint_column(
    mock_server, second_server, tmp_path
) -> None:
    mock_server.config = MockStreamConfig(**FAST)
    base = make_config(mock_server, concurrency=6, total_requests=18, warmup_requests=2)
    endpoints = [
        EndpointSpec("left", mock_server.url, metrics_url=f"{mock_server.url}/metrics", weight=2),
        EndpointSpec("right", second_server.url, metrics_url=f"{second_server.url}/metrics", weight=1),
    ]
    result = await run_fanout(base, endpoints, schedule_dir=tmp_path)
    by_label = {}
    for metric in result.request_metrics:
        by_label.setdefault(metric.endpoint_label, []).append(metric)
    assert {k: len(v) for k, v in by_label.items()} == {"left": 12, "right": 6}
    assert mock_server.max_running == 4 and second_server.max_running == 2
    # one start instant: both endpoints' first requests leave together
    first = [min(m.start_time for m in ms) for ms in by_label.values()]
    assert abs(first[0] - first[1]) < 0.01
    assert all(m.start_time >= result.scenario_start_monotonic for m in result.request_metrics)
    assert result.manifest["concurrency_split"] == {"left": 4, "right": 2}
    summary = result.get_summary()
    assert summary["steady_state"]["available"]
    assert summary["server_agreement"]["left"]["tokens"]["match"]
    assert summary["server_agreement"]["right"]["tokens"]["match"]
    result.export_csv(tmp_path / "merged.csv")
    header = (tmp_path / "merged.csv").read_text().splitlines()[0].split(",")
    assert "endpoint" in header and "fixture_id" in header
    result.export_xlsx(tmp_path / "merged.xlsx")
    assert "Requests" in load_workbook(tmp_path / "merged.xlsx").sheetnames
    assert (tmp_path / "schedule_left.jsonl").exists()


async def test_fixture_id_is_carried_into_every_record(mock_server, tmp_path) -> None:
    handler = FixtureHandler(write_fixtures(tmp_path / "fx.jsonl"))
    mock_server.config = MockStreamConfig(**FAST)
    result = await run_benchmark(
        mock_server, modality_handler=handler, total_requests=4,
        modality_config={"modality": "fixture", "model": "mock-model"},
    )
    assert sorted(m.fixture_id for m in result.request_metrics) == ["fx-a", "fx-a", "fx-b", "fx-b"]
    sent = [r.body["messages"][0]["content"] for r in mock_server.requests]
    assert sorted(sent) == ["alpha", "alpha", "beta", "beta"]
    assert result.manifest["fixture_file_sha256"]
    assert result.manifest["fixture_manifest_sha256"] == handler.manifest_sha256()


async def test_replay_selects_the_fixture_named_in_the_trace(mock_server, tmp_path) -> None:
    handler = FixtureHandler(write_fixtures(tmp_path / "fx.jsonl"))
    trace = tmp_path / "trace.jsonl"
    trace.write_text('{"offset_s": 0, "fixture_id": "fx-b"}\n{"offset_s": 0.02, "fixture_id": "fx-b"}\n')
    mock_server.config = MockStreamConfig(**FAST)
    result = await run_benchmark(
        mock_server, modality_handler=handler, arrival_mode="replay", replay_file=str(trace),
        total_requests=0, modality_config={"modality": "fixture", "model": "mock-model"},
    )
    assert [m.fixture_id for m in result.request_metrics] == ["fx-b", "fx-b"]


async def test_metrics_sidecar_reports_spec_decode_prefix_cache_kv_and_queue(mock_server) -> None:
    mock_server.config = MockStreamConfig(
        output_tokens=12, tokens_per_chunk=3, ttft_s=0.01, token_interval_s=0.004,
        prompt_tokens=10, cached_prompt_tokens=4,
    )
    result = await run_benchmark(
        mock_server, concurrency=2, total_requests=4,
        metrics_url=f"{mock_server.url}/metrics", metrics_scrape_interval_seconds=0.02,
    )
    summary = result.server_metrics["summary"]
    counters, derived, gauges = summary["counters"], summary["derived"], summary["gauges"]
    # 4 requests x 3 follow-up chunks x 2 accepted draft tokens
    assert counters["spec_decode_accepted_tokens"]["delta"] == 24
    assert derived["spec_decode_acceptance_rate"] == pytest.approx(1.0)
    assert derived["spec_decode_mean_acceptance_length"] == pytest.approx(3.0)
    assert derived["prefix_cache_hit_rate"] == pytest.approx(0.4)
    assert {"kv_cache_usage", "waiting_requests", "running_requests"} <= set(gauges)
    assert summary["histograms"]["itl"]["count_delta"] == 12


# -- ladder --------------------------------------------------------------------

def level(c, tput, errors=0.0, ttft=0.5, decode=60.0):
    return {"concurrency": c, "throughput_tok_s": tput, "error_rate_percent": errors,
            "ttft_p95_s": ttft, "decode_p5_tok_s": decode}


def test_ladder_stops_after_two_levels_without_throughput_gain() -> None:
    decision = evaluate_ladder(
        [level(1, 100), level(2, 190), level(4, 360), level(8, 370), level(16, 372)],
        LadderRules(),
    )
    assert decision["stopped"] and decision["stop_rules"] == ["throughput_gain"]
    assert decision["stopped_at_concurrency"] == 16
    assert decision["knee_concurrency"] == 4


def test_ladder_single_bad_level_does_not_stop() -> None:
    decision = evaluate_ladder(
        [level(1, 100), level(2, 200, errors=3.0), level(4, 400), level(8, 800)], LadderRules()
    )
    assert not decision["stopped"]
    assert decision["levels"][1]["flags"] == ["error_rate"]


@pytest.mark.parametrize("bad, rule", [
    (dict(errors=5.0), "error_rate"),
    (dict(ttft=9.0), "ttft_p95"),
    (dict(decode=10.0), "p5_decode"),
])
def test_ladder_each_rule(bad, rule) -> None:
    levels = [level(1, 100), level(2, 200, **bad), level(4, 400, **bad)]
    decision = evaluate_ladder(levels, LadderRules(ttft_p95_limit_s=2.0))
    assert decision["stop_rules"] == [rule] and decision["knee_concurrency"] == 1


# -- soak, burst, config CLI ---------------------------------------------------

async def test_soak_windows(mock_server) -> None:
    mock_server.config = MockStreamConfig(**FAST)
    result = await run_benchmark(mock_server, concurrency=2, duration_seconds=0.6)
    windows = soak_windows(result, window_seconds=0.2)
    full = [w for w in windows if not w["partial"]]
    assert len(full) == 3
    assert all(w["partial"] for w in windows[3:])  # drain after the deadline, if any
    assert sum(w["requests_completed"] for w in windows) == len(result.request_metrics)
    assert all(w["output_tokens_per_second"] > 0 for w in full)


def _lab_toml(tmp_path: Path, urls: list[str], **extra: str) -> dict:
    cfg = {
        "model": "mock-model",
        "modality": "fixture",
        "fixtures": str(write_fixtures(tmp_path / "fx.jsonl")),
        "concurrency": 2,
        "total_requests": 6,
        "warmup_requests": 2,
        "repeats": 2,
        "exports": ["json", "csv", "xlsx"],
        "show_progress": False,
        "slo": {"ttft_p95_s": 1.0, "decode_p5_tok_s": 10, "max_stall_s": 1.0},
        "endpoints": [
            {"label": f"ep{i}", "url": url, "metrics_url": f"{url}/metrics"}
            for i, url in enumerate(urls)
        ],
        "ladder": {"levels": [1, 2, 4], "requests_per_concurrency": 3,
                   "min_p5_decode_tok_s": 1.0},
        "burst": {"cooldown_s": 0},
    }
    cfg.update(extra)
    return cfg


async def test_config_cli_fixed_mode_with_repeats(mock_server, tmp_path) -> None:
    mock_server.config = MockStreamConfig(**FAST)
    summary = await run_from_config(_lab_toml(tmp_path, [mock_server.url]), tmp_path / "out")
    out = Path(summary["output_dir"])
    for rep in ("r1", "r2"):
        for name in ("results.json", "requests.csv", "results.xlsx", "manifest.json"):
            assert (out / rep / name).exists(), (rep, name)
        assert list((out / rep).glob("schedule_*.jsonl"))
    repeats = summary["repeats"]
    assert repeats["repeats"] == 2
    assert repeats["metrics"]["ttft_p50_s"]["cv"] is not None
    assert json.loads((out / "lab_summary.json").read_text())["mode"] == "fixed"
    assert "resolved" in json.loads((out / "plan.json").read_text())


async def test_config_cli_burst_and_ladder(mock_server, tmp_path) -> None:
    mock_server.config = MockStreamConfig(**FAST)
    burst = await run_from_config(
        _lab_toml(tmp_path, [mock_server.url], mode="burst", concurrency=4, repeats=1),
        tmp_path / "out",
    )
    assert burst["repeats"]["per_repeat"][0]["error_rate_percent"] == 0
    assert mock_server.max_running == 4
    ladder = await run_from_config(_lab_toml(tmp_path, [mock_server.url], mode="ladder"),
                                   tmp_path / "out")
    assert [lv["concurrency"] for lv in ladder["ladder"]["levels"]] == [1, 2, 4]


def test_interactive_entry_point_is_kept() -> None:
    from llm_load_tester import __main__ as entry
    from llm_load_tester.cli import run_interactive_config

    assert callable(run_interactive_config)
    import inspect
    assert "run_interactive_config()" in inspect.getsource(entry.main)
